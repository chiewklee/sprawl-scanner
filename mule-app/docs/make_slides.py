"""Two intro slides for the Sprawl Scanner demo: why we built it, and what it does."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# Salesforce (SLDS) palette, per the slides standing rules
INK = RGBColor(0x18, 0x18, 0x18)
INK2 = RGBColor(0x44, 0x44, 0x44)
MUTED = RGBColor(0x74, 0x74, 0x74)
SURFACE = RGBColor(0xF3, 0xF3, 0xF3)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xC9, 0xC9, 0xC9)
ACCENT = RGBColor(0x01, 0x76, 0xD3)        # brand blue
ACCENT_DARK = RGBColor(0x01, 0x44, 0x86)   # dark brand blue
NAVY = RGBColor(0x03, 0x2D, 0x60)
CLOUD = RGBColor(0x1B, 0x96, 0xFF)
SKY = RGBColor(0x00, 0xA1, 0xE0)
GOOD = RGBColor(0x2E, 0x84, 0x4A)
WARN = RGBColor(0xFE, 0x93, 0x39)
SERIOUS = RGBColor(0xEA, 0x00, 0x1E)         # redundant: SLDS error red (gray area keeps warning orange)
CRIT = RGBColor(0xEA, 0x00, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Salesforce Sans"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def dot(slide, x, y, d, fill, glyph, glyph_color=INK):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.name = FONT
    r.font.size = Pt(d * 40)
    r.font.bold = True
    r.font.color.rgb = glyph_color
    return s


def background(slide):
    box(slide, 0, 0, 13.333, 7.5, fill=SURFACE, shape=MSO_SHAPE.RECTANGLE)


def header(slide, kicker, title, subtitle=None):
    box(slide, 0, 0, 13.333, 2.0, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    box(slide, 0, 2.0, 13.333, 0.05, fill=SKY, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.6, 0.42, 12, 0.3, kicker, size=12, color=CLOUD, bold=True)
    text(slide, 0.6, 0.75, 12.2, 0.8, title, size=32, bold=True, color=WHITE)
    if subtitle:
        text(slide, 0.6, 1.5, 12.2, 0.4, subtitle, size=15, color=BORDER)


def footer(slide, left):
    text(slide, 0.6, 7.06, 12, 0.3, left, size=10, color=MUTED)


# ============================================================== slide 1: why
s = prs.slides.add_slide(BLANK)
background(s)
header(s, "WHY WE BUILT THIS",
       "Everyone sells a catalog. Nobody finds the sprawl.",
       "As agents, MCP servers and APIs multiply, the same capability gets built again and again, and nothing flags it.")

# The ask: quote card
box(s, 0.6, 2.35, 6.9, 3.05, fill=CARD, line=BORDER)
box(s, 0.6, 2.35, 0.09, 3.05, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.0, 2.52, 6.3, 0.3, "THE ASK", size=11, color=ACCENT, bold=True)
text(s, 1.0, 2.86, 6.25, 1.9,
     "“As you bring in thousands of agents, MCP servers and APIs, I want to hit a button "
     "that says: here’s where you have uniqueness, here’s where you have redundancy, "
     "here’s where you have gray areas.”",
     size=19, color=INK, spacing=1.2)
text(s, 1.0, 4.82, 6.3, 0.35,
     [[("Fassil Molla, KPMG", {"bold": True, "color": INK}), ("  ·  use-case session, 24 September", {"color": MUTED})]],
     size=13)

# The gap: what exists vs what's missing
box(s, 7.8, 2.35, 4.95, 3.05, fill=CARD, line=BORDER)
text(s, 8.1, 2.52, 4.4, 0.3, "THE GAP", size=11, color=ACCENT, bold=True)
rows = [
    (GOOD, "✓", WHITE, "Registry / catalog", "lists what you have"),
    (GOOD, "✓", WHITE, "Agent evals", "grade how well one agent performs"),
    (CRIT, "✕", WHITE, "Overlap detection", "which capabilities exist more than once?"),
]
for i, (c, g, gc, name, what) in enumerate(rows):
    y = 2.95 + i * 0.76
    dot(s, 8.1, y + 0.03, 0.36, c, g, gc)
    text(s, 8.65, y, 3.9, 0.7,
         [[(name, {"bold": True, "size": 15})], [(what, {"color": INK2, "size": 13})]], spacing=1.05)
text(s, 8.65, 5.0, 3.9, 0.3, "Nobody offers this today.", size=13, color=CRIT, bold=True)

# Why it matters
text(s, 0.6, 5.58, 6, 0.3, "WHY IT MATTERS", size=11, color=ACCENT, bold=True)
tiles = [
    ("Cost", "Every duplicate is built, run, secured and maintained more than once."),
    ("Risk", "Five copies of a capability means five places to patch, govern and audit."),
    ("Agent quality", "Near-identical tools make agents choose inconsistently."),
    ("Speed", "Reuse what exists, and check before anyone writes code."),
]
w = (12.15 - 3 * 0.2) / 4
for i, (t, d) in enumerate(tiles):
    x = 0.6 + i * (w + 0.2)
    box(s, x, 5.88, w, 1.07, fill=CARD, line=BORDER)
    text(s, x + 0.2, 5.97, w - 0.35, 0.95, [[(t, {"bold": True, "size": 14})], [(d, {"color": INK2, "size": 11})]], spacing=1.05)
footer(s, "Sprawl Scanner  ·  credit to Fassil Molla (KPMG) for the ask")

# ========================================================== slide 2: what
s = prs.slides.add_slide(BLANK)
background(s)
header(s, "SO WE BUILT IT",
       "One button. Three answers: unique, redundant, gray.",
       "A working prototype on MuleSoft, running on CloudHub 2.0 next to the assets it scans.")

# Flow: Scan -> Assets -> Sprawl
steps = [
    ("1", "Scan", "Reads every API, MCP server and agent in Exchange, including live tool lists from running MCP servers."),
    ("2", "Assets", "Breaks each asset into what it can do: API operations, MCP tools, agent skills. Flags undocumented assets."),
    ("3", "Sprawl", "Compares every capability with every other, by meaning. Clusters duplicates and recommends what to consolidate."),
]
w = 3.75
for i, (n, t, d) in enumerate(steps):
    x = 0.6 + i * (w + 0.45)
    box(s, x, 2.35, w, 1.5, fill=CARD, line=BORDER)
    dot(s, x + 0.25, 2.52, 0.42, ACCENT, n, WHITE)
    text(s, x + 0.82, 2.55, w - 1.0, 0.4, t, size=19, bold=True)
    text(s, x + 0.25, 3.08, w - 0.45, 1.1, d, size=12.5, color=INK2, spacing=1.12)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.08), Inches(2.94), Inches(0.3), Inches(0.32))
        a.fill.solid()
        a.fill.fore_color.rgb = BORDER
        a.line.fill.background()

# How it decides
box(s, 0.6, 4.05, 12.15, 1.15, fill=CARD, line=BORDER)
text(s, 0.85, 4.2, 2.2, 0.3, "HOW IT DECIDES", size=11, color=ACCENT, bold=True)
how = [
    ("Meaning, not words", "Gemini embeddings: “look up customer” matches “get account”."),
    ("Judge for close calls", "A Gemini model rules duplicate, overlap or distinct, with a reason."),
    ("Knows good patterns", "MCP-enabling an API, or an agent calling its tools, is expected, not sprawl."),
]
cw = (12.15 - 0.5) / 3
for i, (t, d) in enumerate(how):
    x = 0.85 + i * cw
    text(s, x, 4.5, cw - 0.3, 0.65, [[(t, {"bold": True, "size": 13.5})], [(d, {"color": INK2, "size": 11.5})]], spacing=1.05)

# Results on our Exchange
text(s, 0.6, 5.45, 8, 0.3, "FIRST RUN ON OUR EXCHANGE", size=11, color=ACCENT, bold=True)
kpis = [
    ("76", "assets", None, None),
    ("197", "capabilities", None, None),
    ("73", "redundant", SERIOUS, "⧉"),
    ("49", "gray area", WARN, "◐"),
    ("75", "unique", GOOD, "✓"),
    ("11", "undocumented", MUTED, "!"),
]
w = (12.15 - 5 * 0.18) / 6
for i, (v, label, c, g) in enumerate(kpis):
    x = 0.6 + i * (w + 0.18)
    box(s, x, 5.78, w, 1.1, fill=CARD, line=BORDER)
    text(s, x + 0.2, 5.84, w - 0.3, 0.55, v, size=28, bold=True)
    lx = x + 0.2
    if c is not None:
        dot(s, lx, 6.5, 0.22, c, g, WHITE if c in (CRIT, GOOD, SERIOUS, MUTED) else INK)
        lx += 0.3
    text(s, lx, 6.47, w - (lx - x) - 0.05, 0.3, label, size=12, color=INK2)
footer(s, "e.g. four copies of the same golf MCP server; the PTO capability on two MCP servers; two identical SE Insights brokers"
          "  ·  credit: Fassil Molla (KPMG)")

import pathlib
out = str(pathlib.Path(__file__).with_name("sprawl-scanner-intro.pptx"))
prs.save(out)
print("saved", out)
